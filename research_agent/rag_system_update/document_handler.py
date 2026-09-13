"""Document loading, lightweight metadata extraction, and chunking."""

import csv
import json
import re
from pathlib import Path
from typing import Any

from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import HTMLHeaderTextSplitter, RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pptx import Presentation

from ..config import settings

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".html",
    ".htm",
    ".py",
    ".xml",
    ".xlsx",
    ".pptx",
}


def _clean_text(text: str) -> str:
    """Remove layout-only whitespace without changing the actual wording."""
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _find_heading(document: Document) -> str:
    """Return the most useful heading already present in loader metadata/text."""
    for key, value in document.metadata.items():
        if value and ("header" in key.lower() or "heading" in key.lower()):
            return str(value).strip()

    for line in document.page_content.splitlines():
        line = line.strip()
        if line.startswith("#"):
            heading = line.lstrip("# ").strip()
            if heading:
                return heading
    return ""


def _make_metadata(document: Document, filename: str) -> dict[str, str]:
    """Keep only metadata useful to a research answer."""
    title = (
        document.metadata.get("title")
        or document.metadata.get("document_title")
        or Path(filename).stem.replace("_", " ")
    )
    return {
        "filename": Path(filename).name,
        "title": str(title).strip(),
        "heading": _find_heading(document),
    }


def _load_text_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()

    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        return [Document(page_content=json.dumps(payload, ensure_ascii=False, indent=2))]

    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = csv.reader(handle)
            text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        return [Document(page_content=text)]

    if suffix in {".html", ".htm"}:
        splitter = HTMLHeaderTextSplitter(
            headers_to_split_on=[
                ("h1", "Header 1"),
                ("h2", "Header 2"),
                ("h3", "Header 3"),
            ]
        )
        return splitter.split_text(path.read_text(encoding="utf-8", errors="ignore"))

    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        documents: list[Document] = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(" | ".join(str(value or "") for value in row))
            documents.append(Document(page_content="\n".join(rows)))
        return documents

    if suffix == ".pptx":
        presentation = Presentation(path)
        return [
            Document(
                page_content="\n".join(
                    shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
                )
            )
            for slide in presentation.slides
        ]

    return TextLoader(str(path), encoding="utf-8", autodetect_encoding=True).load()


def _load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported document type: {suffix}. Supported types: {supported}")

    if suffix == ".pdf":
        return PyPDFLoader(str(path), extract_images=False).load()
    if suffix == ".docx":
        return Docx2txtLoader(str(path)).load()
    return _load_text_file(path)


class DocumentHandler:
    """Own document files and turn them into retrieval-ready chunks."""

    def __init__(self, storage_dir: str | Path):
        self.storage_dir = Path(storage_dir)
        self.storage_dir.mkdir(parents=True, exist_ok=True)

    def save_upload(self, filename: str, content: bytes) -> Path:
        """Save an upload unless a file with the same name is already present."""
        destination = self.storage_dir / Path(filename).name
        if destination.exists():
            print(f"[RAG][FILES] Already exists, keeping existing file: {destination.name}")
            return destination

        destination.write_bytes(content)
        print(f"[RAG][FILES] Saved: {destination.name}")
        return destination

    def file_exists(self, filename: str) -> bool:
        return (self.storage_dir / Path(filename).name).exists()

    def prepare_file(self, path: str | Path, filename: str | None = None) -> list[Document]:
        """Load, clean, annotate, and split one file."""
        file_path = Path(path)
        source_name = filename or file_path.name
        print(f"[RAG][DOCUMENT] Reading: {source_name}")

        source_documents = _load_file(file_path)
        prepared: list[Document] = []

        for document in source_documents:
            content = _clean_text(document.page_content)
            if not content:
                continue
            prepared.append(
                Document(
                    page_content=content,
                    metadata=_make_metadata(document, source_name),
                )
            )

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=getattr(settings, "rag_chunk_size", 900),
            chunk_overlap=getattr(settings, "rag_chunk_overlap", 140),
            separators=["\n\n", "\n", ". ", "; ", ", ", " ", ""],
            add_start_index=False,
        )
        chunks = splitter.split_documents(prepared)

        for chunk in chunks:
            chunk.page_content = _clean_text(chunk.page_content)

        print(f"[RAG][CHUNKING] Created {len(chunks)} chunks from {source_name}")
        return [chunk for chunk in chunks if chunk.page_content]

    def read_stored_file(self, filename: str, max_chars: int = 12000) -> str:
        """Read the original stored file for non-RAG file inspection."""
        path = self.storage_dir / Path(filename).name
        if not path.exists():
            return f"File not found: {filename}"

        documents = _load_file(path)
        return "\n\n".join(document.page_content for document in documents)[:max_chars]
