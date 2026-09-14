"""Generic document handler logic without file-type-specific metadata extraction."""

import re
from pathlib import Path

from langchain_community.document_loaders import Docx2txtLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation

from ...config import settings

SUPPORTED_EXTENSIONS = {".pdf", ".md", ".ppt", ".pptx", ".docx", ".txt"}


def _clean_text(text: str) -> str:
    """Normalize whitespace and remove stray null characters from raw text."""
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


class DocumentHandler:
    """Base document loader and chunker for supported research document types."""

    supported_extensions = SUPPORTED_EXTENSIONS

    def __init__(self, storage_directory: str | Path):
        """Create the session storage directory used for uploaded files."""
        self.storage_dir = Path(storage_directory)
        self.storage_dir.mkdir(parents=True, exist_ok=True)

    def save_upload(self, uploaded_file_path: str | Path, file_contents: bytes) -> Path:
        """Persist a file in the session folder, replacing an older copy with the same name."""
        original_path = Path(uploaded_file_path)
        if original_path.is_dir():
            raise ValueError("save_upload expects a file path, not a directory.")

        destination = self.storage_dir / original_path.name
        if destination.exists():
            destination.unlink()
            print(f"[RAG][FILES] Replaced existing file: {destination.name}")

        destination.write_bytes(file_contents)
        print(f"[RAG][FILES] Saved: {destination.name}")
        return destination

    def remove_file(self, stored_file_path: str | Path) -> bool:
        """Delete the stored copy of a file from the session folder."""
        file_path = Path(stored_file_path)
        destination = self.storage_dir / file_path.name
        if not destination.exists():
            return False

        destination.unlink()
        print(f"[RAG][FILES] Removed: {destination.name}")
        return True

    def file_exists(self, stored_file_path: str | Path) -> bool:
        """Return True when a file with the same basename exists in the session storage folder."""
        return (self.storage_dir / Path(stored_file_path).name).exists()

    def _get_document_handler(self, document_path: str | Path):
        """Return the PDF-specific handler for PDF files and the base handler for other supported types."""
        path = Path(document_path)
        if path.suffix.lower() == ".pdf":
            from .pdf_handler import PDFDocumentHandler
            return PDFDocumentHandler(self.storage_dir)
        return self

    def load_document(self, document_path: Path) -> list[Document]:
        """Load a supported document from disk into a list of LangChain documents."""
        suffix = document_path.suffix.lower()
        if suffix not in self.supported_extensions:
            supported = ", ".join(sorted(self.supported_extensions))
            raise ValueError(f"Unsupported document type: {suffix}. Supported types: {supported}")

        if suffix in {".md", ".txt"}:
            return TextLoader(str(document_path), encoding="utf-8", autodetect_encoding=True).load()

        if suffix == ".docx":
            return Docx2txtLoader(str(document_path)).load()

        if suffix in {".ppt", ".pptx"}:
            presentation = Presentation(str(document_path))
            return [
                Document(
                    page_content="\n".join(
                        shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
                    )
                )
                for slide in presentation.slides
            ]

        supported = ", ".join(sorted(self.supported_extensions))
        raise ValueError(f"Unsupported document type: {suffix}. Supported types: {supported}")

    def extract_metadata(self, document_path: str | Path) -> dict[str, str]:
        """Build the default metadata block for a source document before chunking."""
        return {
            "filename": Path(document_path).name,
            "source": Path(document_path).name,
        }

    def generate_chunks(self, source_documents: list[Document], source_name: str) -> list[Document]:
        """Split a list of loaded documents into smaller chunks and clean their text."""
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=getattr(settings, "rag_chunk_size", 900),
            chunk_overlap=getattr(settings, "rag_chunk_overlap", 140),
            separators=["\n\n", "\n", ". ", "; ", ", ", " ", ""],
            add_start_index=False,
        )
        chunks = splitter.split_documents(source_documents)
        for chunk in chunks:
            chunk.page_content = _clean_text(chunk.page_content)

        print(f"[RAG][CHUNKING] Created {len(chunks)} chunks from {source_name}")
        return [chunk for chunk in chunks if chunk.page_content]

    def prepare_document(self, document_path: Path) -> list[Document]:
        """Load one document, clean its content, and prepare chunked output for indexing."""
        source_name = document_path.name
        print(f"[RAG][DOCUMENT] Reading: {source_name}")

        source_documents = self.load_document(document_path)
        prepared_documents: list[Document] = []

        for source_document in source_documents:
            content = _clean_text(source_document.page_content)
            if not content:
                continue
            prepared_documents.append(
                Document(
                    page_content=content,
                    metadata=self.extract_metadata(document_path),
                )
            )

        return self.generate_chunks(prepared_documents, source_name)

    def prepare_file(self, input_file_path: str | Path) -> list[Document]:
        """Validate a single file path, delegate to the proper handler, and return chunked documents."""
        file_path = Path(input_file_path)
        if file_path.is_dir():
            raise ValueError(
                "DocumentHandler.prepare_file expects a single file path, not a directory. "
                "Directory iteration should happen in the test entrypoint."
            )
        if not file_path.exists():
            raise FileNotFoundError(f"Input file does not exist: {file_path}")

        handler = self._get_document_handler(file_path)
        return handler.prepare_document(file_path)

    def read_stored_file(self, stored_file_name: str | Path, max_chars: int = 12000) -> str:
        """Read the original stored file for non-RAG inspection and return a trimmed preview."""
        file_path = self.storage_dir / Path(stored_file_name).name
        if not file_path.exists():
            return f"File not found: {stored_file_name}"

        handler = self._get_document_handler(file_path)
        documents = handler.load_document(file_path)
        return "\n\n".join(document.page_content for document in documents)[:max_chars]
