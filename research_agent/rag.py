import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pptx import Presentation
from langchain_chroma import Chroma
from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import (
    HTMLHeaderTextSplitter,
    RecursiveCharacterTextSplitter,
)

from .config import settings

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".html",
    ".htm",
    ".py",
    ".xml",
    ".xlsx",
    ".pptx",
}


def _stable_id(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


class HybridRAG:
    """Persistent document pipeline with semantic, lexical, and reciprocal-rank retrieval."""

    def __init__(self, session_id: str):
        self.session_id = session_id
        self.session_documents_dir = Path(settings.documents_dir) / session_id
        self.session_chroma_dir = Path(settings.chroma_dir) / session_id
        self.session_documents_dir.mkdir(parents=True, exist_ok=True)
        self.session_chroma_dir.mkdir(parents=True, exist_ok=True)
        self.embeddings = HuggingFaceEmbeddings(
            model_name=settings.embedding_model_id,
            cache_folder=settings.embedding_cache_dir,
            model_kwargs={
                "device": "cpu",
                "local_files_only": settings.embedding_local_files_only,
            },
            encode_kwargs={"normalize_embeddings": True},
        )
        self.vectorstore = Chroma(
            collection_name=f"research_documents_{session_id.replace('-', '_')}",
            embedding_function=self.embeddings,
            persist_directory=str(self.session_chroma_dir),
        )
        self.documents: dict[str, Document] = {}
        self.registry_path = self.session_documents_dir / "document_registry.json"
        self.files: dict[str, dict[str, Any]] = self._load_registry()
        self._load_existing_chunks()

    def _load_registry(self) -> dict[str, dict[str, Any]]:
        if not self.registry_path.exists():
            return {}
        try:
            return json.loads(self.registry_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return {}

    def _load_existing_chunks(self) -> None:
        existing = self.vectorstore.get(include=["documents", "metadatas"])
        for content, metadata in zip(
            existing.get("documents", []), existing.get("metadatas", [])
        ):
            document = Document(page_content=content, metadata=metadata or {})
            self.documents[document.metadata.get("chunk_id", _stable_id(content))] = (
                document
            )

    def _save_registry(self) -> None:
        self.registry_path.parent.mkdir(parents=True, exist_ok=True)
        self.registry_path.write_text(
            json.dumps(self.files, indent=2), encoding="utf-8"
        )

    def add_documents(
        self,
        documents: list[Document],
        source_name: str,
        file_metadata: dict[str, Any] | None = None,
    ) -> int:
        """Normalize, enrich, split, deduplicate, and persist a document collection."""
        if not documents:
            return 0
        source_id = _stable_id(f"{self.session_id}:{source_name.lower()}")[:16]
        normalized: list[Document] = []
        for page_number, document in enumerate(documents, start=1):
            content = _normalize_text(document.page_content)
            if not content:
                continue
            metadata = dict(document.metadata)
            metadata.update(
                {
                    "source": source_name,
                    "session_id": self.session_id,
                    "source_id": source_id,
                    **(file_metadata or {}),
                    "page": metadata.get("page", page_number),
                }
            )
            normalized.append(Document(page_content=content, metadata=metadata))

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.rag_chunk_size,
            chunk_overlap=settings.rag_chunk_overlap,
            separators=["\n\n", "\n", ". ", "; ", ", ", " ", ""],
            add_start_index=True,
        )
        chunks = splitter.split_documents(normalized)
        unique_chunks: list[Document] = []
        ids: list[str] = []
        for chunk_index, chunk in enumerate(chunks):
            chunk.page_content = _normalize_text(chunk.page_content)
            chunk_hash = _stable_id(f"{source_id}:{chunk.page_content}")
            chunk.metadata.update(
                {
                    "chunk_id": chunk_hash,
                    "chunk_index": chunk_index,
                    "content_hash": chunk_hash,
                }
            )
            if chunk_hash in self.documents:
                continue
            unique_chunks.append(chunk)
            ids.append(chunk_hash)
            self.documents[chunk_hash] = chunk
        if unique_chunks:
            self.vectorstore.add_documents(unique_chunks, ids=ids)
        self.files[source_name] = {
            "source": source_name,
            "session_id": self.session_id,
            "source_id": source_id,
            "file_type": Path(source_name).suffix.lower(),
            **(file_metadata or {}),
            "chunk_count": len(unique_chunks),
            "pages": len(normalized),
        }
        self._save_registry()
        return len(unique_chunks)

    def save_uploaded_file(self, file_name: str, content: bytes) -> Path:
        """Persist the original upload inside this session's document directory."""
        safe_name = Path(file_name).name
        destination = self.session_documents_dir / safe_name
        destination.write_bytes(content)
        return destination

    def list_files(self) -> list[dict[str, Any]]:
        return list(self.files.values())

    def read_file(
        self, source_name: str, query: str = "", max_chars: int = 12000
    ) -> str:
        """Read a stored file or the most relevant excerpts from it."""
        matches = [
            doc
            for doc in self.documents.values()
            if doc.metadata.get("source") == source_name
        ]
        if not matches:
            available = ", ".join(self.files) or "none"
            return f"File not found: {source_name}. Available stored files: {available}"
        if query:
            terms = set(re.findall(r"\w+", query.lower()))
            matches.sort(
                key=lambda doc: len(
                    terms & set(re.findall(r"\w+", doc.page_content.lower()))
                ),
                reverse=True,
            )
        text = "\n\n".join(
            f"[{source_name} | page {doc.metadata.get('page', '?')} | chunk {doc.metadata.get('chunk_index', '?')}]\n{doc.page_content}"
            for doc in matches
        )
        return text[:max_chars]

    def retrieve(self, query: str, k: int = 8) -> list[dict[str, Any]]:
        """Fuse vector and lexical rankings, then apply diversity-aware source ranking."""
        if not self.documents:
            return []
        dense_docs = self.vectorstore.similarity_search_with_relevance_scores(
            query, k=min(k * 3, 30)
        )
        terms = set(re.findall(r"\w+", query.lower()))
        lexical = sorted(
            self.documents.values(),
            key=lambda doc: len(
                terms & set(re.findall(r"\w+", doc.page_content.lower()))
            )
            / max(len(terms), 1),
            reverse=True,
        )[: min(k * 3, 30)]
        rankings: dict[str, dict[str, Any]] = {}
        for rank, (doc, dense_score) in enumerate(dense_docs, start=1):
            chunk_id = doc.metadata.get("chunk_id", _stable_id(doc.page_content))
            rankings.setdefault(
                chunk_id,
                {
                    "document": doc,
                    "dense_score": float(dense_score),
                    "lexical_score": 0.0,
                    "rrf": 0.0,
                },
            )
            rankings[chunk_id]["rrf"] += 1 / (60 + rank)
        for rank, doc in enumerate(lexical, start=1):
            chunk_id = doc.metadata.get("chunk_id", _stable_id(doc.page_content))
            rankings.setdefault(
                chunk_id,
                {"document": doc, "dense_score": 0.0, "lexical_score": 0.0, "rrf": 0.0},
            )
            rankings[chunk_id]["lexical_score"] = len(
                terms & set(re.findall(r"\w+", doc.page_content.lower()))
            ) / max(len(terms), 1)
            rankings[chunk_id]["rrf"] += 1 / (60 + rank)
        selected: list[dict[str, Any]] = []
        per_source: dict[str, int] = {}
        for item in sorted(
            rankings.values(), key=lambda value: value["rrf"], reverse=True
        ):
            document = item["document"]
            source = document.metadata.get("source", "uploaded document")
            if per_source.get(source, 0) >= 3:
                continue
            per_source[source] = per_source.get(source, 0) + 1
            selected.append(
                {
                    "content": document.page_content,
                    "source": source,
                    "page": document.metadata.get("page", "?"),
                    "chunk_id": document.metadata.get("chunk_id", ""),
                    "score": round(item["rrf"], 6),
                    "citation": f"{source}, page {document.metadata.get('page', '?')}",
                }
            )
            if len(selected) >= k:
                break
        return selected


def _load_structured_text(path: str) -> list[Document]:
    suffix = Path(path).suffix.lower()
    if suffix == ".json":
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
        return [
            Document(
                page_content=json.dumps(payload, indent=2), metadata={"format": "json"}
            )
        ]
    if suffix == ".csv":
        with Path(path).open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        return [
            Document(
                page_content=json.dumps(row, ensure_ascii=False),
                metadata={"row": index + 1, "format": "csv"},
            )
            for index, row in enumerate(rows)
        ]
    if suffix in {".html", ".htm"}:
        text_splitter = HTMLHeaderTextSplitter(
            headers_to_split_on=[
                ("h1", "Header 1"),
                ("h2", "Header 2"),
                ("h3", "Header 3"),
            ]
        )
        return text_splitter.split_text(
            Path(path).read_text(encoding="utf-8", errors="ignore")
        )
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        return [
            Document(
                page_content="\n".join(
                    " | ".join(str(value or "") for value in row)
                    for row in sheet.iter_rows(values_only=True)
                ),
                metadata={"sheet": sheet.title, "format": "xlsx"},
            )
            for sheet in workbook.worksheets
        ]
    if suffix == ".pptx":
        presentation = Presentation(path)
        return [
            Document(
                page_content="\n".join(
                    shape.text for shape in slide.shapes if hasattr(shape, "text")
                ),
                metadata={"slide": index + 1, "format": "pptx"},
            )
            for index, slide in enumerate(presentation.slides)
        ]
    return TextLoader(path, encoding="utf-8", autodetect_encoding=True).load()


def load_uploaded_file(path: str) -> list[Document]:
    suffix = Path(path).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported document type: {suffix}. Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if suffix == ".pdf":
        return PyPDFLoader(path, extract_images=False).load()
    if suffix == ".docx":
        return Docx2txtLoader(path).load()
    return _load_structured_text(path)
