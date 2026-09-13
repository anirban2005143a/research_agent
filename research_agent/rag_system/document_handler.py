"""Uploaded-file reading, storage, metadata, and chunk preparation."""

import csv
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pptx import Presentation
from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import HTMLHeaderTextSplitter, RecursiveCharacterTextSplitter

from ..config import settings

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm", ".py", ".xml", ".xlsx", ".pptx"}


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _tokens(text: str) -> set[str]:
    return set(re.findall(r"[\w-]+", text.lower()))


def _document_metadata(document: Document, source_name: str) -> dict[str, Any]:
    metadata = dict(document.metadata)
    headings: list[str] = []
    for key, value in metadata.items():
        if "header" in key.lower() or "heading" in key.lower():
            if value and str(value) not in headings:
                headings.append(str(value))
    for line in document.page_content.splitlines():
        candidate = line.strip()
        if candidate.startswith("#"):
            heading = candidate.lstrip("# ").strip()
            if heading and heading not in headings:
                headings.append(heading)
    title = metadata.get("title") or metadata.get("document_title") or Path(source_name).stem.replace("_", " ")
    metadata.update({"filename": Path(source_name).name, "title": str(title), "headings": " | ".join(headings)})
    return metadata


def _load_structured_text(path: str) -> list[Document]:
    suffix = Path(path).suffix.lower()
    if suffix == ".json":
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
        return [Document(page_content=json.dumps(payload, indent=2), metadata={"format": "json"})]
    if suffix == ".csv":
        with Path(path).open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        return [Document(page_content=json.dumps(row, ensure_ascii=False), metadata={"row": index + 1, "format": "csv"}) for index, row in enumerate(rows)]
    if suffix in {".html", ".htm"}:
        splitter = HTMLHeaderTextSplitter(headers_to_split_on=[("h1", "Header 1"), ("h2", "Header 2"), ("h3", "Header 3")])
        return splitter.split_text(Path(path).read_text(encoding="utf-8", errors="ignore"))
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        return [Document(page_content="\n".join(" | ".join(str(value or "") for value in row) for row in sheet.iter_rows(values_only=True)), metadata={"sheet": sheet.title, "format": "xlsx"}) for sheet in workbook.worksheets]
    if suffix == ".pptx":
        presentation = Presentation(path)
        return [Document(page_content="\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")), metadata={"slide": index + 1, "format": "pptx"}) for index, slide in enumerate(presentation.slides)]
    return TextLoader(path, encoding="utf-8", autodetect_encoding=True).load()


def _load_uploaded_file(path: str) -> list[Document]:
    """Load one supported uploaded file into page-like documents."""
    suffix = Path(path).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported document type: {suffix}. Supported types: {supported}")
    if suffix == ".pdf":
        return PyPDFLoader(path, extract_images=False).load()
    if suffix == ".docx":
        return Docx2txtLoader(path).load()
    return _load_structured_text(path)


class DocumentHandler:
    """Manage original uploaded files and prepare them for vector storage."""

    def __init__(self, storage_dir: str | Path, session_id: str | None = None):
        self.storage_dir = Path(storage_dir)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        self.session_id = session_id or self.storage_dir.name
        self.registry_path = self.storage_dir / "document_registry.json"
        self.files: dict[str, dict[str, Any]] = self._load_registry()

    def _load_registry(self) -> dict[str, dict[str, Any]]:
        if not self.registry_path.exists():
            return {}
        try:
            return json.loads(self.registry_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return {}

    def _save_registry(self) -> None:
        self.registry_path.write_text(json.dumps(self.files, indent=2), encoding="utf-8")

    def read_file(self, path: str | Path) -> list[Document]:
        return _load_uploaded_file(str(path))

    def save_upload(self, file_name: str, content: bytes) -> Path:
        destination = self.storage_dir / Path(file_name).name
        destination.write_bytes(content)
        return destination

    def prepare_file(self, path: str | Path, source_name: str | None = None, file_metadata: dict[str, Any] | None = None) -> list[Document]:
        """Read, normalize, enrich, and chunk one file for vector storage."""
        file_path = Path(path)
        source = source_name or file_path.name
        prepared: list[Document] = []
        for page_number, document in enumerate(self.read_file(file_path), start=1):
            content = _normalize_text(document.page_content)
            if not content:
                continue
            metadata = _document_metadata(document, source)
            metadata.update({"source": source, "session_id": self.session_id, **(file_metadata or {}), "page": metadata.get("page", page_number)})
            prepared.append(Document(page_content=content, metadata=metadata))
        splitter = RecursiveCharacterTextSplitter(chunk_size=settings.rag_chunk_size, chunk_overlap=settings.rag_chunk_overlap, separators=["\n\n", "\n", ". ", "; ", ", ", " ", ""], add_start_index=True)
        chunks = splitter.split_documents(prepared)
        for chunk in chunks:
            chunk.page_content = _normalize_text(chunk.page_content)
        return chunks

    def register_file(self, source_name: str, file_metadata: dict[str, Any] | None = None, pages: int = 0, chunk_count: int = 0) -> None:
        self.files[source_name] = {"source": source_name, "session_id": self.session_id, "file_type": Path(source_name).suffix.lower(), **(file_metadata or {}), "pages": pages, "chunk_count": chunk_count}
        self._save_registry()

    def list_files(self) -> list[dict[str, Any]]:
        return sorted(self.files.values(), key=lambda item: str(item.get("source", "")).casefold())

    def read_stored_file(self, source_name: str, query: str = "", max_chars: int = 12000) -> str:
        """Read the original stored file, optionally ordering pages by query overlap."""
        path = self.storage_dir / Path(source_name).name
        if not path.exists():
            available = ", ".join(item["source"] for item in self.list_files()) or "none"
            return f"File not found: {source_name}. Available stored files: {available}"
        documents = self.read_file(path)
        if query:
            query_terms = _tokens(query)
            documents.sort(key=lambda doc: len(query_terms & _tokens(doc.page_content)), reverse=True)
        return "\n\n".join(f"[{source_name} | page {document.metadata.get('page', index + 1)}]\n{document.page_content}" for index, document in enumerate(documents))[:max_chars]
